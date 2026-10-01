"""C508 labs/: rewrite the WSQ course identity in the Markdown and the lab DOCX files.

Each Lab-NN-*.pdf is LibreOffice output of the Lab-NN-*.docx beside it, so the PDFs
are re-exported from the fixed DOCX by build_all.sh (not edited in place).
"""
import re
from pathlib import Path
from docx import Document

OLD, NEW, CODE = "Architecture Drawing with Revit", "Autodesk Revit Architecture Masterclass", "C508"

SUBS = [
    ("**WSQ Course Code:** TGS-2021004287", f"**Course Code:** {CODE}"),
    (" · Version v12", " · Version v1.0"),
    (f"{OLD} (TGS-2021004287)", f"{NEW} ({CODE})"),
    (f"# {OLD} — Hands-On Labs", f"# {NEW} — Hands-On Labs"),
    ("TGS-2021004287", CODE),
    # practice set: no assessment in a non-WSQ course
    ("does not replace assessed lab evidence.", "does not replace the 12 course labs."),
    ("extends the 12 assessed labs", "extends the 12 course labs"),
    ("does not replace the metric workflow or evidence required in Labs 1-12.",
     "does not replace the metric workflow practised in Labs 1-12."),
    ("for the assessed course project.", "for the course project."),
]


def text(s):
    for a, b in SUBS:
        s = s.replace(a, b)
    return re.sub(r"\((LO\d) · A\d(?:, A\d)*\)", r"(\1)", s)


def set_par(par, t):
    runs = par.runs
    k = next((i for i, r in enumerate(runs) if r.text), 0)
    for i, r in enumerate(runs):
        r.text = t if i == k else ""


def fix_md():
    n = 0
    for f in Path("labs").rglob("*.md"):
        old = f.read_text(encoding="utf-8")
        new = text(old)
        if new != old:
            f.write_text(new, encoding="utf-8")
            n += 1
    return n


def fix_docx():
    n = 0
    for f in Path("labs").rglob("*.docx"):
        d = Document(str(f))
        changed = False
        parts = [d] + [x for s in d.sections for x in (s.header, s.footer)]
        pars = [p for part in parts for p in part.paragraphs]
        pars += [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
        for p in pars:
            new = text(p.text)
            if new != p.text:
                set_par(p, new)
                changed = True
        if changed:
            d.core_properties.title = text(d.core_properties.title or "")
            d.save(str(f))
            n += 1
    return n


if __name__ == "__main__":
    print(f"md rewritten: {fix_md()}, docx rewritten: {fix_docx()}")
