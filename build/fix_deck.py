"""C508 deck: residual WSQ-layer edits the converter does not cover.

Slide indexes are given in SOURCE-deck terms and mapped through the dropped set.
"""
import re
from pptx import Presentation

F = "courseware/Autodesk Revit Architecture Masterclass (C508)-v1.0.pptx"
DROPPED = {2, 9, 10, 11, 12, 13, 14, 120, 121}
p = Presentation(F)
S = p.slides


def n(old):
    assert old not in DROPPED
    return old - sum(1 for d in DROPPED if d < old)


def set_para(par, text):
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def drop_para(par):
    par._p.getparent().remove(par._p)


def shape(old, sid):
    return next(sh for sh in S[n(old)].shapes if sh.shape_id == sid)


def edit(old, sid, fn):
    tf = shape(old, sid).text_frame
    for par in list(tf.paragraphs):
        new = fn(par.text)
        if new is None:
            drop_para(par)
        elif new != par.text:
            set_para(par, new)


def card(old, sid, head, body):
    pars = shape(old, sid).text_frame.paragraphs
    assert len(pars) == 2, (old, sid, len(pars))
    set_para(pars[0], head)
    set_para(pars[1], body)


# trainer profile
edit(4, 21, lambda t: t.replace("WSQ courses on", "Courses on"))
# Download Course Material — no assessment in a non-WSQ course
SUBS7 = {
    "•  Keep them handy — the final assessment is open book (slides + Learner Guide).":
        "•  Keep them handy — refer to the slides and Learner Guide during the labs.",
    "•  Your certificate is also issued on this portal after you are assessed Competent.":
        "•  Your certificate of completion is also issued on this portal after the course.",
}
edit(7, 7, lambda t: SUBS7.get(t, t))
# schedule slide
edit(8, 5, lambda t: t.replace("8 hours/day", "7.5 hours/day"))
SUBS8 = {
    "–  Digital attendance (AM) · introductions": "–  Welcome · introductions · course overview",
    "–  Digital attendance (PM) · Day 1 recap": "–  Day 1 recap & Q&A",
}
edit(8, 11, lambda t: SUBS8.get(t, t))


def day2(t):
    if t.startswith("–  Final Assessment"):
        return None
    t = t.replace("Documentation, Specification & Assessment", "Documentation & Specification")
    t = t.replace("–  TRAQOM survey · assessment attendance", "–  Course summary & feedback")
    t = t.replace("–  9:30am–6:30pm · 1-hour lunch · tea breaks within", "–  9:30am–5:30pm · 30-min lunch")
    return t


edit(8, 12, day2)
# Key Concepts eyebrows carried the TSC A#/K# mapping
for old, k in ((27, 1), (52, 2), (80, 3), (101, 4)):
    edit(old, 4, lambda t, k=k: f"TOPIC 0{k} · KEY CONCEPTS" if re.match(r"(TSC|C508) MAPPING", t) else t)
# lab cover cards: keep the LO, drop the TSC ability codes
for i, s in enumerate(S):
    for sh in s.shapes:
        if sh.has_text_frame:
            for par in sh.text_frame.paragraphs:
                t = par.text
                new = re.sub(r"\((LO\d) · A\d(?:, A\d)*\)", r"(\1)", t)
                if new != t:
                    set_para(par, new)
edit(115, 4, lambda t: "COURSE RECAP" if t == "LEARNING OUTCOMES DELIVERED" else t)
# recommended courses: no programme labels
for sid in (11, 16, 21, 26, 31):
    edit(117, sid, lambda t: t.replace(" (WSQ).", "."))
edit(117, 36, lambda t: t.replace("— WSQ-funded CAD/BIM pathway.", "— the full CAD/BIM pathway."))
# support slide card 4
card(118, 26, "Certificate",
     "Download your certificate of completion from https://lms-tms.tertiaryinfotech.com")
# WSQ Assessment slide -> Course Feedback (rewritten in place, card design kept)
edit(119, 5, lambda t: "Course Feedback" if t == "Assessment" else t)
card(119, 11, "Share your feedback",
     "Complete the course feedback form on the LMS before you leave.")
card(119, 16, "Help us improve",
     "Tell us what worked well and what we should change for the next class.")
card(119, 21, "Certificate of completion",
     "Issued on the LMS after the course — download it from the portal.")
card(119, 26, "Course portal",
     "Slides, Learner Guide and labs: https://lms-tms.tertiaryinfotech.com")
p.save(F)
print("fix_deck ok:", len(S), "slides")
